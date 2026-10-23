"""Text extraction from uploaded documents (.txt, .md, .docx, .pdf)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass


class ExtractError(ValueError):
    pass


@dataclass
class Extracted:
    text: str
    kind: str
    pages: int | None = None


_MAX_PDF_PAGES = 60


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes, *, max_chars: int) -> Extracted:
    name = (filename or "").lower()
    if not data:
        raise ExtractError("The file is empty.")

    if name.endswith((".txt", ".md", ".markdown", ".csv", ".log")):
        out = Extracted(text=_decode(data), kind="text")
    elif name.endswith(".docx"):
        out = _extract_docx(data)
    elif name.endswith(".pdf"):
        out = _extract_pdf(data)
    elif _looks_like_text(data):
        out = Extracted(text=_decode(data), kind="text")
    else:
        raise ExtractError("Unsupported file type. Upload .txt, .md, .docx or .pdf.")

    out.text = _normalise(out.text)
    if not out.text:
        raise ExtractError("No readable text was found in the file.")
    if len(out.text) > max_chars:
        out.text = out.text[:max_chars]
    return out


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _looks_like_text(data: bytes) -> bool:
    sample = data[:4096]
    if b"\x00" in sample:
        return False
    printable = sum(1 for b in sample if 32 <= b < 127 or b in (9, 10, 13) or b >= 128)
    return printable / max(1, len(sample)) > 0.9


def _extract_docx(data: bytes) -> Extracted:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("DOCX support is not installed.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("Could not read the Word document.") from exc

    parts: list[str] = []
    for para in document.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            parts.append(f"\n## {t}")
        elif "list" in style:
            parts.append(f"- {t}")
        else:
            parts.append(t)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return Extracted(text="\n".join(parts), kind="docx")


def _extract_pdf(data: bytes) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("PDF support is not installed.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractError("The PDF is password protected.") from exc
        pages = len(reader.pages)
        chunks = []
        for page in reader.pages[:_MAX_PDF_PAGES]:
            chunks.append(page.extract_text() or "")
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("Could not read the PDF.") from exc
    return Extracted(text="\n\n".join(chunks), kind="pdf", pages=pages)
