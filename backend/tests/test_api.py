from __future__ import annotations

import io
import json

import pytest

from app.services.llm import LLMResult

MODEL_JSON = {
    "company_name": "TechNova",
    "newsletter_type": "Monthly Update",
    "language": "en",
    "hero_title": "AI Assistant 2.0 Lands",
    "subtitle": "A big month.",
    "summary": "TechNova shipped a lot. " * 10,
    "date": "October 2026",
    "estimated_read_time": "3 min read",
    "highlights": [
        {"icon": "🚀", "title": "AI Assistant 2.0", "subtitle": "Shipped"},
        {"icon": "📈", "title": "100k users", "subtitle": "Milestone"},
    ],
    "sections": [
        {"icon": "🚀", "title": "Product", "content": "**AI Assistant 2.0** shipped."},
        {"icon": "👥", "title": "People", "content": "We are hiring."},
    ],
    "stats": [{"value": "100k", "label": "users"}],
    "quote": None,
    "cta": {"label": "Read more", "url": "https://technova.example"},
    "image_keywords": ["modern office", "software engineers"],
    "footer": {"tagline": "Build the future"},
}


@pytest.fixture
def fake_llm(monkeypatch):
    calls = []

    async def complete_json(self, system, user, **kw):
        calls.append((system, user, kw))
        return LLMResult(content=json.dumps(MODEL_JSON), provider="groq", model="fake", usage={"total_tokens": 10})

    monkeypatch.setattr("app.services.llm.LLMService.complete_json", complete_json)
    return calls


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok"
    assert "groq" in body["providers"]
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert "X-Request-Id" in r.headers


def test_templates_and_presets(client):
    assert {t["id"] for t in client.get("/api/templates").json()} >= {"modern", "editorial", "minimal", "bold"}
    presets = client.get("/api/brands/presets").json()
    avl = next(p for p in presets if p["id"] == "avl")
    assert avl["logo"].startswith("data:image/svg+xml;base64,")
    assert avl["primary"] == "#23a638"


def test_generate_end_to_end(client, fake_llm):
    payload = {
        "sources": [{"kind": "text", "text": "- Launched AI Assistant 2.0\n- 100k users"}],
        "options": {"highlight_count": 2, "section_count": 2},
        "brand": {"company_name": "TechNova", "primary": "#111111", "secondary": "#222222", "accent": "#333333"},
        "render": {"template": "bold", "theme": "light"},
        "auto_hero": "generated",
    }
    r = client.post("/api/generate", json=payload)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["newsletter"]["hero_title"] == "AI Assistant 2.0 Lands"
    assert body["newsletter"]["hero_image"]["kind"] == "generated"
    assert body["html"].startswith("<!DOCTYPE html>")
    assert "AI Assistant 2.0 Lands" in body["html"]
    assert body["model"] == "fake"
    # Prompt-injection guard: sources are wrapped and labelled as untrusted.
    system, user, _ = fake_llm[0]
    assert "<<<SOURCE 1" in user and "untrusted" in user.lower()


def test_generate_validation_errors(client, fake_llm):
    r = client.post("/api/generate", json={"sources": []})
    assert r.status_code == 422
    r = client.post("/api/generate", json={"sources": [{"text": "   "}]})
    assert r.status_code == 422
    r = client.post("/api/generate", json={"sources": [{"text": "x"}], "brand": {"primary": "red"}})
    assert r.status_code == 422


def test_generate_repairs_bad_json(client, monkeypatch):
    responses = iter(["Here you go: {not json", json.dumps(MODEL_JSON)])

    async def complete_json(self, system, user, **kw):
        return LLMResult(content=next(responses), provider="groq", model="fake")

    monkeypatch.setattr("app.services.llm.LLMService.complete_json", complete_json)
    r = client.post("/api/generate", json={"sources": [{"text": "update"}], "auto_hero": "none"})
    assert r.status_code == 200, r.text
    assert any("repaired" in w for w in r.json()["warnings"])


def test_render_and_export(client, sample_newsletter, sample_brand):
    payload = {"newsletter": sample_newsletter.model_dump(), "brand": sample_brand.model_dump(), "render": {"template": "editorial", "theme": "dark"}}
    r = client.post("/api/render", json=payload)
    assert r.status_code == 200 and "TechNova" in r.json()["html"]
    for fmt, needle in (("html", "<!DOCTYPE"), ("email", "<!DOCTYPE"), ("markdown", "# AI Assistant"), ("json", '"hero_title"'), ("text", "AI Assistant")):
        r = client.post("/api/export", json={**payload, "format": fmt})
        assert r.status_code == 200, fmt
        assert needle in r.text
        assert "attachment" in r.headers["content-disposition"]


def test_upload_svg_logo_is_sanitised(client):
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" onload="alert(1)"><script>alert(1)</script><rect width="1" height="1"/></svg>'
    r = client.post("/api/upload/image?kind=logo", files={"file": ("logo.svg", svg, "image/svg+xml")})
    assert r.status_code == 200, r.text
    import base64

    decoded = base64.b64decode(r.json()["data_uri"].split(",", 1)[1]).decode()
    assert "script" not in decoded and "onload" not in decoded


def test_upload_rejects_svg_hero_and_bad_types(client):
    r = client.post("/api/upload/image?kind=hero", files={"file": ("x.svg", b"<svg xmlns='http://www.w3.org/2000/svg'/>", "image/svg+xml")})
    assert r.status_code == 415
    r = client.post("/api/upload/image?kind=hero", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 415
    r = client.post("/api/upload/image?kind=hero", files={"file": ("x.png", b"not a png", "image/png")})
    assert r.status_code == 422


def test_extract_txt_and_docx(client):
    r = client.post("/api/extract", files={"file": ("notes.txt", b"hello\r\n\r\n\r\nworld", "text/plain")})
    assert r.status_code == 200 and r.json()["text"] == "hello\n\nworld"

    import docx

    doc = docx.Document()
    doc.add_heading("Launch", level=1)
    doc.add_paragraph("We shipped it.")
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post("/api/extract", files={"file": ("notes.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200
    assert "## Launch" in r.json()["text"] and "We shipped it." in r.json()["text"]


@pytest.mark.parametrize(
    "url",
    [
        "http://localhost/admin",
        "http://127.0.0.1/",
        "http://169.254.169.254/latest/meta-data/",
        "http://10.0.0.1/",
        "ftp://example.com/x",
        "http://user:pw@example.com/",
        "http://example.com:8080/",
    ],
)
def test_fetch_url_blocks_ssrf(client, url):
    r = client.post("/api/fetch-url", json={"url": url})
    assert r.status_code == 400, r.text


def test_access_key_gate(monkeypatch, settings):
    from fastapi.testclient import TestClient

    from app.main import create_app

    monkeypatch.setattr(settings, "app_access_key", "secret")
    with TestClient(create_app(settings)) as c:
        assert c.get("/api/health").status_code == 200
        assert c.get("/api/templates").status_code == 200  # meta routes stay open
        r = c.post("/api/render", json={"newsletter": {}})
        assert r.status_code == 401
        r = c.post("/api/render", json={"newsletter": {"sections": [{"title": "a", "content": "b"}]}}, headers={"X-App-Key": "secret"})
        assert r.status_code == 200


def test_body_size_limit(client, settings):
    big = "x" * (settings.max_json_body_bytes + 10)
    r = client.post("/api/render", content=big, headers={"Content-Type": "application/json"})
    assert r.status_code == 413
